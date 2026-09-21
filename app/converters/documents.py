import fitz
from docx import Document


def write_document(text, target, output):
    if target == 'txt':
        output.write_text(text, encoding='utf-8')
    elif target == 'docx':
        doc = Document()
        for line in text.splitlines():
            doc.add_paragraph(line)
        doc.save(output)
    elif target == 'pdf':
        # Story performs flowing pagination and supports Unicode through MuPDF fonts.
        from html import escape
        html = '<html><body>' + ''.join('<p>' + escape(line or ' ') + '</p>' for line in text.splitlines()) + '</body></html>'
        story = fitz.Story(html=html, user_css='body { font-family: sans-serif; font-size: 11pt; } p { margin: 0 0 5pt; white-space: pre-wrap; }')
        writer = fitz.DocumentWriter(str(output))
        page_rect = fitz.paper_rect('a4')
        content_rect = page_rect + (45, 45, -45, -45)
        more = True
        while more:
            device = writer.begin_page(page_rect)
            more, _ = story.place(content_rect)
            story.draw(device)
            writer.end_page()
        writer.close()


def convert_document(source, output, options, progress):
    progress(10, 'Reading document')
    if source.suffix == '.docx':
        doc = Document(source)
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        parts = []
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            elif isinstance(block, Table):
                parts.extend('\t'.join(cell.text for cell in row.cells) for row in block.rows)
        text = '\n'.join(parts)
    else:
        text = source.read_text(encoding='utf-8-sig')
    write_document(text, output.suffix[1:], output)
    return output
