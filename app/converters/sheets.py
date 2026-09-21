import csv
import json
from datetime import date, datetime
from openpyxl import Workbook, load_workbook


def safe_cell(value):
    # Spreadsheet apps interpret these prefixes as executable formulas.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value


def write_rows(rows, target, output):
    if target == 'csv':
        with output.open('w', newline='', encoding='utf-8-sig') as stream:
            csv.writer(stream).writerows([[safe_cell(v) for v in row] for row in rows])
    elif target == 'xlsx':
        book = Workbook(write_only=True)
        sheet = book.create_sheet('Converted')
        for row in rows:
            sheet.append([safe_cell(v) for v in row])
        book.save(output)
    elif target == 'json':
        output.write_text(json.dumps(rows, ensure_ascii=False, indent=2, default=lambda x: x.isoformat() if isinstance(x, (date, datetime)) else str(x)), encoding='utf-8')


def read_rows(source, extension):
    if extension == 'csv':
        with source.open(encoding='utf-8-sig', newline='') as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
            except csv.Error:
                dialect = csv.excel
            return list(csv.reader(stream, dialect))
    if extension == 'xlsx':
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            # The UI explicitly states that CSV/JSON export uses the first sheet.
            sheet = book.worksheets[0]
            if sheet.max_row and sheet.max_row > 200000:
                raise ValueError('This spreadsheet exceeds the 200,000 row limit.')
            if sheet.max_column and sheet.max_column > 1000:
                raise ValueError('This spreadsheet exceeds the 1,000 column limit.')
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            book.close()
    data = json.loads(source.read_text(encoding='utf-8-sig'))
    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list):
        raise ValueError('JSON must contain an array of objects, arrays, or values.')
    if data and all(isinstance(item, dict) for item in data):
        headers = list(dict.fromkeys(key for item in data for key in item))
        return [headers] + [[json.dumps(item.get(key), ensure_ascii=False) if isinstance(item.get(key), (list, dict)) else item.get(key, '') for key in headers] for item in data]
    return [[json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in (item if isinstance(item, list) else [item])] for item in data]


def convert_sheet(source, output, options, progress):
    progress(10, 'Reading spreadsheet')
    rows = read_rows(source, source.suffix[1:])
    write_rows(rows, output.suffix[1:], output)
    return output
