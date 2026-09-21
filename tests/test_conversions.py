import csv
import io
import json
import shutil
import time
import zipfile
from pathlib import Path
import fitz
import pytest
from PIL import Image, ImageDraw, ImageFont
from fastapi.testclient import TestClient
from openpyxl import load_workbook, Workbook
from docx import Document
from app.main import app, RATE
from app.converters import CONVERTERS
from app.converters.pdf import page_selection
from app.services.validation import inspect_file
from app.services.jobs import FILES, JOBS, cleanup


@pytest.fixture
def client():
    RATE.clear()
    with TestClient(app) as api:
        yield api
    for record in FILES.values():
        shutil.rmtree(record['path'].parent, ignore_errors=True)
    FILES.clear()
    JOBS.clear()


def pdf_bytes(lines=None, pages=1, password=None):
    doc = fitz.open()
    for number in range(pages):
        page = doc.new_page()
        page.insert_text((50, 70), '\n'.join(lines or [f'Page {number + 1}', 'WAYNE ENTERPRISES', 'Invoice Number: 2381', 'Price: KSh 137,999', 'Thank you.']))
    kwargs = dict(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password) if password else {}
    result = doc.tobytes(**kwargs)
    doc.close()
    return result


def upload(client, name, contents, mime):
    response = client.post('/api/files', files={'file': (name, contents, mime)})
    assert response.status_code == 201, response.text
    return response.json()


def run(client, record, target, **options):
    response = client.post('/api/conversions', json={'file_id': record['id'], 'output_format': target, **options})
    assert response.status_code == 202, response.text
    job_id = response.json()['conversion_id']
    for _ in range(300):
        status = client.get('/api/conversions/' + job_id).json()
        if status['status'] != 'processing':
            return status
        time.sleep(.1)
    pytest.fail('Conversion did not finish within 30 seconds')


def result(client, job):
    assert job['status'] == 'completed', job
    response = client.get(job['download_url'])
    assert response.status_code == 200
    return response


def test_raw_never_invents_columns(client):
    source = upload(client, 'invoice.pdf', pdf_bytes(), 'application/pdf')
    response = result(client, run(client, source, 'csv'))
    rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert rows[0] == ['text']
    assert all(len(row) == 1 for row in rows)
    assert ['Price: KSh 137,999'] in rows
    assert response.headers['content-type'].startswith('text/csv')
    response = result(client, run(client, source, 'xlsx'))
    book = load_workbook(io.BytesIO(response.content))
    assert book.active.max_column == 1
    assert book.active['A2'].value == 'WAYNE ENTERPRISES'


def test_page_selection_and_images(client):
    source = upload(client, 'report.pdf', pdf_bytes(pages=3), 'application/pdf')
    assert source['pages'] == 3
    text = result(client, run(client, source, 'txt', pages='2')).text
    assert 'Page 2' in text and 'Page 1' not in text
    archive = result(client, run(client, source, 'png', pages='1,3'))
    with zipfile.ZipFile(io.BytesIO(archive.content)) as images:
        assert images.namelist() == ['page-001.png', 'page-003.png']
    single = result(client, run(client, source, 'jpg', pages='2'))
    assert single.headers['content-type'] == 'image/jpeg'
    assert Image.open(io.BytesIO(single.content)).format == 'JPEG'


def test_layout_and_tables(client):
    doc = fitz.open()
    page = doc.new_page()
    for x in (40, 220, 350):
        page.draw_line((x, 40), (x, 130))
    for y in (40, 70, 100, 130):
        page.draw_line((40, y), (350, y))
    for row, cells in enumerate([['Product', 'Price'], ['Milk', '400'], ['Bread', '120']]):
        for col, text in enumerate(cells):
            page.insert_text((50 + 180 * col, 60 + row * 30), text)
    source = upload(client, 'table.pdf', doc.tobytes(), 'application/pdf')
    for mode in ('layout', 'table'):
        rows = list(csv.reader(io.StringIO(result(client, run(client, source, 'csv', mode=mode)).content.decode('utf-8-sig'))))
        assert rows[1] == ['Milk', '400']
    paragraph = upload(client, 'paragraph.pdf', pdf_bytes(), 'application/pdf')
    failure = run(client, paragraph, 'csv', mode='table')
    assert failure['status'] == 'failed' and 'No tables' in failure['error']


def test_password_and_invalid_pages(client):
    source = upload(client, 'protected.pdf', pdf_bytes(password='secret'), 'application/pdf')
    assert source['encrypted']
    failure = run(client, source, 'txt', password='wrong')
    assert 'password protected' in failure['error']
    assert 'WAYNE' in result(client, run(client, source, 'txt', password='secret')).text
    failure = run(client, source, 'txt', password='secret', pages='2-5')
    assert 'valid pages' in failure['error']


def test_validation_and_deletion(client):
    for name, contents, mime in [('fake.pdf', b'not a pdf', 'application/pdf'), ('bad.exe', b'MZ', 'application/octet-stream'), ('empty.csv', b'', 'text/csv'), ('bad.png', b'not png', 'image/png'), ('bad.xlsx', b'not zip', 'application/octet-stream'), ('bad.json', b'{', 'application/json'), ('binary.txt', b'abc\x00', 'text/plain')]:
        assert client.post('/api/files', files={'file': (name, contents, mime)}).status_code in (415, 422)
    source = upload(client, '../../data.csv', b'name,value\nWayne,100\n', 'text/csv')
    assert source['name'] == 'data.csv'
    job = run(client, source, 'xlsx')
    assert result(client, job).content.startswith(b'PK')
    assert client.delete('/api/files/' + source['id']).status_code == 204
    assert client.get(job['download_url']).status_code == 404


def test_spreadsheet_json_and_formula_safety(client):
    source = upload(client, 'input.json', b'[{"name":"Wayne","note":"=1+1"}]', 'application/json')
    data = result(client, run(client, source, 'csv')).content.decode('utf-8-sig')
    assert "'=1+1" in data
    xlsx = result(client, run(client, source, 'xlsx')).content
    book = load_workbook(io.BytesIO(xlsx))
    assert book.active['B2'].data_type == 's'
    source2 = upload(client, 'book.xlsx', xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert json.loads(result(client, run(client, source2, 'json')).content)[1][0] == 'Wayne'


def test_documents_and_image_roundtrip(client):
    source = upload(client, 'hello.txt', b'Hello File Fix\nSecond line', 'text/plain')
    pdf = result(client, run(client, source, 'pdf')).content
    with fitz.open(stream=pdf, filetype='pdf') as doc:
        assert 'Hello File Fix' in doc[0].get_text()
    word = result(client, run(client, source, 'docx')).content
    source2 = upload(client, 'hello.docx', word, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    assert 'Second line' in result(client, run(client, source2, 'txt')).text
    stream = io.BytesIO()
    Image.new('RGBA', (100, 100), (0, 0, 0, 0)).save(stream, 'PNG')
    source3 = upload(client, 'image.png', stream.getvalue(), 'image/png')
    jpg = result(client, run(client, source3, 'jpg'))
    assert Image.open(io.BytesIO(jpg.content)).getpixel((50, 50)) == (255, 255, 255)
    assert result(client, run(client, source3, 'pdf')).content.startswith(b'%PDF')


def test_scan_detection_and_ocr(client):
    image = Image.new('RGB', (1200, 350), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((60, 80), 'FILE FIX SCANNED DOCUMENT', fill='black', font=ImageFont.load_default(size=48))
    png = io.BytesIO(); image.save(png, 'PNG')
    doc = fitz.open(); page = doc.new_page(width=600, height=175)
    page.insert_image(page.rect, stream=png.getvalue())
    source = upload(client, 'scan.pdf', doc.tobytes(), 'application/pdf')
    assert source['likely_scanned']
    assert 'Enable OCR' in run(client, source, 'txt')['error']
    if shutil.which('tesseract'):
        assert 'FILE FIX' in result(client, run(client, source, 'txt', ocr=True)).text.upper()


def test_expiry(client):
    source = upload(client, 'expire.txt', b'expire me', 'text/plain')
    path = FILES[source['id']]['path']
    FILES[source['id']]['expires'] = time.time() - 1
    cleanup()
    assert not path.exists()
    assert source['id'] not in FILES


@pytest.mark.parametrize('expression', ['0', '3-1', '1--2', 'abc', '1,', '9999'])
def test_bad_page_ranges(expression):
    with pytest.raises(ValueError):
        page_selection(expression, 10)


@pytest.mark.parametrize('source_format,target', sorted(CONVERTERS))
def test_every_registered_conversion(tmp_path, source_format, target):
    source = tmp_path / ('source.' + source_format)
    output = tmp_path / ('output.' + target)
    if source_format == 'pdf':
        source.write_bytes(pdf_bytes())
    elif source_format in ('png', 'jpg', 'webp'):
        Image.new('RGB', (60, 40), 'white').save(source)
    elif source_format == 'csv':
        source.write_text('name,value\nFile Fix,123\n', encoding='utf-8')
    elif source_format == 'json':
        source.write_text('[{"name":"File Fix","value":123}]', encoding='utf-8')
    elif source_format == 'xlsx':
        book = Workbook(); book.active.append(['name', 'value']); book.active.append(['File Fix', 123]); book.save(source)
    elif source_format == 'docx':
        doc = Document(); doc.add_paragraph('File Fix document'); doc.save(source)
    else:
        source.write_text('File Fix document', encoding='utf-8')
    actual = CONVERTERS[(source_format, target)](source, output, {}, lambda *args: None)
    assert actual.exists() and actual.stat().st_size > 0
    if target in ('pdf', 'png', 'jpg', 'docx', 'xlsx', 'txt', 'json', 'csv'):
        inspect_file(actual, target)
    if target == 'html':
        assert '<!doctype html>' in actual.read_text()


def test_upload_page_and_size_limits(client, monkeypatch):
    import app.main as main
    monkeypatch.setattr(main, 'MAX_BYTES', 10)
    response = client.post('/api/files', files={'file': ('too-big.txt', b'x' * 11, 'text/plain')})
    assert response.status_code == 422 and 'larger' in response.json()['detail']
    monkeypatch.setattr(main, 'MAX_BYTES', 25 * 1024 * 1024)
    import app.services.validation as validation
    monkeypatch.setattr(validation, 'MAX_PAGES', 2)
    response = client.post('/api/files', files={'file': ('too-many.pdf', pdf_bytes(pages=3), 'application/pdf')})
    assert response.status_code == 422 and 'page limit' in response.json()['detail']


def test_blank_pdf_and_html_escaping(client):
    doc = fitz.open(); doc.new_page()
    source = upload(client, 'blank.pdf', doc.tobytes(), 'application/pdf')
    assert 'readable text' in run(client, source, 'txt')['error']
    source = upload(client, 'html.pdf', pdf_bytes(lines=['<script>alert(1)</script>']), 'application/pdf')
    html = result(client, run(client, source, 'html')).text
    assert '<script>' not in html and '&lt;script&gt;' in html
